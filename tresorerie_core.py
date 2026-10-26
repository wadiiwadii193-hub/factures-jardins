"""
tresorerie_core.py — logique métier (sans Streamlit) pour le suivi de trésorerie
des jardins familiaux.

  * détection automatique de la structure du tableau (mois, blocs Entrées / Sorties)
  * ajout d'un montant dans la cellule (catégorie × mois) avec sauvegarde de sécurité
  * lecture du tableau avec recalcul des totaux (openpyxl ne calcule pas les formules)
  * numérisation de document (OpenCV), OCR (Tesseract / EasyOCR), extraction
    du montant TTC, de la date et suggestion de catégorie
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import os
import re
import shutil
import unicodedata
from copy import copy
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
import pandas as pd
from PIL import Image, ImageOps

MOIS_FR = ["janvier", "février", "mars", "avril", "mai", "juin", "juillet",
           "août", "septembre", "octobre", "novembre", "décembre"]


def norm(s) -> str:
    """Minuscule, sans accents, espaces compactés : sert à comparer des libellés."""
    s = "" if s is None else str(s)
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", s).strip().lower()


MOIS_NORM = [norm(m) for m in MOIS_FR]


class LayoutError(ValueError):
    """La structure attendue du tableau n'a pas été trouvée."""


class OcrUnavailable(RuntimeError):
    """Le moteur OCR demandé n'est pas utilisable (non installé, etc.)."""


# ════════════════════════════════════════════════════════════════════════════
#  1. STRUCTURE DU TABLEAU EXCEL
# ════════════════════════════════════════════════════════════════════════════
@dataclass
class Block:
    title_row: int
    first_row: int      # première ligne de catégorie
    last_row: int       # dernière ligne de catégorie (juste avant « Total … »)
    total_row: int

    def rows(self):
        return range(self.first_row, self.last_row + 1)


@dataclass
class Layout:
    label_col: int
    header_row: int
    month_cols: dict            # {1: n° colonne de janvier, ..., 12: décembre}
    start_col: int | None       # « Début d'activité »
    total_col: int | None       # « TOTAL »
    entrees: Block
    sorties: Block
    year: int | None


@dataclass
class Row:
    label: str
    row: int


def _find_label(ws, target: str, cols=(1, 2, 3), min_row=1):
    for c in cols:
        for r in range(min_row, ws.max_row + 1):
            if norm(ws.cell(r, c).value) == target:
                return r, c
    return None


def _block(ws, title: str, total_prefix: str, col: int, row: int) -> Block:
    total_row = None
    for r in range(row + 1, ws.max_row + 1):
        if norm(ws.cell(r, col).value).startswith(total_prefix):
            total_row = r
            break
    if total_row is None:
        raise LayoutError(f"Ligne « {total_prefix} » introuvable sous « {title} ».")
    return Block(title_row=row, first_row=row + 1, last_row=total_row - 1, total_row=total_row)


def detect_layout(ws) -> Layout:
    # ligne des mois
    header_row, month_cols = None, {}
    for r in range(1, min(ws.max_row, 25) + 1):
        found = {}
        for c in range(1, ws.max_column + 1):
            n = norm(ws.cell(r, c).value)
            if n in MOIS_NORM:
                found[MOIS_NORM.index(n) + 1] = c
        if len(found) >= 6:
            header_row, month_cols = r, found
            break
    if header_row is None:
        raise LayoutError("Ligne des mois (janvier … décembre) introuvable.")

    s = _find_label(ws, "sorties")
    if s is None:
        raise LayoutError("Section « Sorties » introuvable (colonne A ou B).")
    s_row, label_col = s
    e = _find_label(ws, "entrees", cols=(label_col,))
    if e is None:
        raise LayoutError("Section « Entrées » introuvable.")
    e_row = e[0]

    start_col = total_col = None
    for c in range(1, ws.max_column + 1):
        n = norm(ws.cell(header_row, c).value)
        if n.startswith("debut"):
            start_col = c
        elif n == "total":
            total_col = c

    year = None
    m = re.search(r"(20\d{2})", ws.title)
    if m:
        year = int(m.group(1))

    return Layout(
        label_col=label_col, header_row=header_row, month_cols=month_cols,
        start_col=start_col, total_col=total_col,
        entrees=_block(ws, "Entrées", "total entrees", label_col, e_row),
        sorties=_block(ws, "Sorties", "total sorties", label_col, s_row),
        year=year,
    )


def valid_sheets(path) -> list[str]:
    """Feuilles dont la structure est reconnue."""
    wb = openpyxl.load_workbook(path)
    out = []
    for ws in wb.worksheets:
        try:
            detect_layout(ws)
            out.append(ws.title)
        except LayoutError:
            pass
    return out


def num(v) -> float:
    """Valeur de cellule → float (les formules du type =5+10 sont additionnées)."""
    if v is None or v == "" or isinstance(v, bool):
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if s.startswith("="):
        expr = s[1:].replace(",", ".")
        if re.fullmatch(r"[\d.+\- ]+", expr):
            return sum(float(t) for t in re.findall(r"[+-]?\d+(?:\.\d+)?", expr))
        return 0.0
    try:
        return float(s.replace("\u00a0", "").replace(" ", "").replace(",", "."))
    except ValueError:
        return 0.0


def _data_cols(ws, lay: Layout) -> list[int]:
    cols = [lay.month_cols[m] for m in range(1, 13)]
    return ([lay.start_col] if lay.start_col else []) + cols


def list_rows(ws, lay: Layout, block: Block):
    """(catégories nommées, lignes libres) d'un bloc."""
    cats, free = [], []
    cols = _data_cols(ws, lay)
    for r in block.rows():
        label = str(ws.cell(r, lay.label_col).value or "").strip()
        if label:
            cats.append(Row(label, r))
        elif all(ws.cell(r, c).value in (None, "") for c in cols):
            free.append(r)
    return cats, free


# ════════════════════════════════════════════════════════════════════════════
#  2. LECTURE DU TABLEAU (avec totaux recalculés)
# ════════════════════════════════════════════════════════════════════════════
@dataclass
class Model:
    ws: object
    layout: Layout
    sorties: list[Row]
    sorties_free: list[int]
    entrees: list[Row]
    entrees_free: list[int]
    month_headers: list[str]
    entrees_df: pd.DataFrame = field(default=None)
    sorties_df: pd.DataFrame = field(default=None)
    summary_df: pd.DataFrame = field(default=None)

    def value(self, row: int, month: int) -> float:
        return num(self.ws.cell(row, self.layout.month_cols[month]).value)


def _block_frame(ws, lay: Layout, block: Block, headers: list[str]) -> pd.DataFrame:
    cols = _data_cols(ws, lay)
    labels, data = [], []
    for r in block.rows():
        label = str(ws.cell(r, lay.label_col).value or "").strip()
        vals = [num(ws.cell(r, c).value) for c in cols]
        if label or any(vals):
            labels.append(label or f"(ligne {r})")
            data.append(vals)
    df = pd.DataFrame(data, index=labels, columns=headers, dtype=float)
    df["Total"] = df.sum(axis=1)
    return df


def open_model(path, sheet: str) -> Model:
    wb = openpyxl.load_workbook(path)
    ws = wb[sheet]
    lay = detect_layout(ws)
    s, s_free = list_rows(ws, lay, lay.sorties)
    e, e_free = list_rows(ws, lay, lay.entrees)

    headers = ([str(ws.cell(lay.header_row, lay.start_col).value).strip()] if lay.start_col else []) \
        + [m.capitalize() for m in MOIS_FR]
    ent = _block_frame(ws, lay, lay.entrees, headers)
    sor = _block_frame(ws, lay, lay.sorties, headers)

    tot_e, tot_s = ent[headers].sum(), sor[headers].sum()
    diff = tot_e - tot_s
    treso = diff.cumsum()
    summary = pd.DataFrame([tot_e, tot_s, diff, treso],
                           index=["Total entrées", "Total sorties", "Différence", "Trésorerie (cumul)"])
    summary["Total"] = [tot_e.sum(), tot_s.sum(), diff.sum(), treso.iloc[-1]]

    return Model(ws=ws, layout=lay, sorties=s, sorties_free=s_free, entrees=e, entrees_free=e_free,
                 month_headers=headers, entrees_df=ent, sorties_df=sor, summary_df=summary)


# ════════════════════════════════════════════════════════════════════════════
#  3. ÉCRITURE DANS EXCEL
# ════════════════════════════════════════════════════════════════════════════
def add_amount(cell, amount: float):
    """Additionne `amount` au contenu de la cellule. Retourne (ancienne, nouvelle valeur)."""
    amt = round(float(amount), 2)
    old = cell.value
    if old in (None, ""):
        new = amt
    elif isinstance(old, bool):
        raise ValueError(f"Cellule {cell.coordinate} : valeur booléenne inattendue.")
    elif isinstance(old, (int, float)):
        new = round(old + amt, 2)
    elif isinstance(old, str) and old.startswith("="):
        new = f"{old}+{amt}"                      # on prolonge la formule existante
    else:
        try:
            new = round(float(str(old).replace(" ", "").replace(",", ".")) + amt, 2)
        except ValueError:
            raise ValueError(f"La cellule {cell.coordinate} contient du texte non numérique : {old!r}")
    cell.value = new
    return old, new


def save_with_backup(wb, path, backup_dir="sauvegardes") -> Path | None:
    path = Path(path)
    backup = None
    if path.exists():
        bdir = path.parent / backup_dir
        bdir.mkdir(exist_ok=True)
        backup = bdir / f"{path.stem}_{dt.datetime.now():%Y%m%d_%H%M%S}{path.suffix}"
        shutil.copy2(path, backup)
    wb.save(path)
    return backup


def _copy_row_style(ws, lay: Layout, src_row: int, dst_row: int):
    cols = [lay.label_col] + _data_cols(ws, lay) + ([lay.total_col] if lay.total_col else [])
    for c in cols:
        ws.cell(dst_row, c)._style = copy(ws.cell(src_row, c)._style)


def record_amount(path, sheet: str, which: str, month: int, amount: float, *,
                  label: str | None = None, new_label: str | None = None,
                  make_backup: bool = True) -> dict:
    """
    Ajoute `amount` dans (catégorie × mois) du bloc `which` ('sorties' ou 'entrees').
    `new_label` crée la catégorie dans la première ligne libre du bloc.
    """
    if not 1 <= int(month) <= 12:
        raise ValueError("Mois invalide.")
    if amount is None or amount <= 0:
        raise ValueError("Le montant doit être strictement positif.")

    wb = openpyxl.load_workbook(path)          # formules conservées
    ws = wb[sheet]
    lay = detect_layout(ws)
    block = lay.sorties if which == "sorties" else lay.entrees
    cats, free = list_rows(ws, lay, block)

    if new_label and new_label.strip():
        known = {norm(c.label): c.row for c in cats}
        if norm(new_label) in known:
            row = known[norm(new_label)]
        else:
            if not free:
                raise LayoutError("Aucune ligne libre dans ce bloc pour créer une catégorie.")
            row = free[0]
            ref = [c.row for c in cats if c.row < row]
            if ref:
                _copy_row_style(ws, lay, ref[-1], row)
            ws.cell(row, lay.label_col).value = new_label.strip()
    else:
        row = next((c.row for c in cats if norm(c.label) == norm(label)), None)
        if row is None:
            raise LayoutError(f"Catégorie « {label} » introuvable dans le bloc {which}.")

    col = lay.month_cols[int(month)]
    cell = ws.cell(row, col)
    old, new = add_amount(cell, amount)
    if make_backup:
        backup = save_with_backup(wb, path)
    else:                     # lot de plusieurs lignes : une seule sauvegarde pour tout le lot
        wb.save(path)
        backup = None
    return {"row": row, "col": col, "cell": cell.coordinate, "old": old, "new": new,
            "label": str(ws.cell(row, lay.label_col).value).strip(), "backup": backup}


def log_operation(path, **fields):
    """Journal CSV (séparateur « ; », lisible dans Excel) pour garder une trace des saisies."""
    log = Path(path).parent / "journal_saisies.csv"
    new = not log.exists()
    with open(log, "a", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        if new:
            w.writerow(["horodatage", *fields.keys()])
        w.writerow([dt.datetime.now().isoformat(timespec="seconds"), *fields.values()])


# ════════════════════════════════════════════════════════════════════════════
#  4. IMAGES : chargement, numérisation, OCR
# ════════════════════════════════════════════════════════════════════════════
def load_document(data: bytes, filename: str = "", max_pages: int = 5):
    """
    Retourne (aperçu PIL, texte intégré au PDF ou None, [images à OCRiser]).
    Un PDF « numérique » est lu directement (pas d'OCR, donc fiable à 100 %).
    """
    is_pdf = filename.lower().endswith(".pdf") or data[:5] == b"%PDF-"
    if not is_pdf:
        img = ImageOps.exif_transpose(Image.open(io.BytesIO(data))).convert("RGB")
        return img, None, [img]

    try:
        import fitz  # PyMuPDF
    except ImportError as e:
        raise OcrUnavailable("Lecture des PDF : installez PyMuPDF (pip install pymupdf).") from e

    doc = fitz.open(stream=data, filetype="pdf")
    pages = [p for i, p in enumerate(doc) if i < max_pages]

    def render(page, dpi):
        pix = page.get_pixmap(matrix=fitz.Matrix(dpi / 72, dpi / 72))
        return Image.frombytes("RGB", (pix.width, pix.height), pix.samples)

    text = "\n".join(p.get_text("text") for p in pages).strip()
    if len(text) >= 40:                       # PDF avec couche texte
        return render(pages[0], 110), text, []
    images = [render(p, 250) for p in pages]  # PDF scanné → OCR
    return images[0], None, images


def _order_points(pts):
    import numpy as np
    rect = np.zeros((4, 2), dtype="float32")
    s, d = pts.sum(axis=1), np.diff(pts, axis=1)
    rect[0], rect[2] = pts[s.argmin()], pts[s.argmax()]
    rect[1], rect[3] = pts[d.argmin()], pts[d.argmax()]
    return rect


def _page_quad(bgr):
    """
    Les 4 coins de la feuille, ou None. Le papier est clair et peu coloré, alors que la table, le bois
    ou le tissu autour sont plus sombres ou plus colorés : on isole le papier par sa couleur,
    ce qui marche beaucoup mieux que la recherche de contours sur un fond texturé.
    """
    import cv2
    import numpy as np

    h, w = bgr.shape[:2]
    k = min(1.0, 900.0 / max(h, w))
    small = cv2.resize(bgr, None, fx=k, fy=k, interpolation=cv2.INTER_AREA) if k < 1 else bgr.copy()
    hsv = cv2.cvtColor(small, cv2.COLOR_BGR2HSV)
    v = cv2.GaussianBlur(hsv[:, :, 2], (0, 0), 4)
    sat = cv2.GaussianBlur(hsv[:, :, 1], (0, 0), 4)
    seuil, _ = cv2.threshold(v, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)
    mask = ((v > seuil * 0.9) & (sat < 95)).astype(np.uint8) * 255
    taille = max(5, int(0.03 * max(small.shape[:2])) | 1)
    mask = cv2.morphologyEx(mask, cv2.MORPH_CLOSE, np.ones((taille, taille), np.uint8))
    mask = cv2.morphologyEx(mask, cv2.MORPH_OPEN, np.ones((taille // 2 | 1, taille // 2 | 1), np.uint8))
    cnts, _ = cv2.findContours(mask, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
    if not cnts:
        return None
    grand = max(cnts, key=cv2.contourArea)
    if cv2.contourArea(grand) < 0.20 * small.shape[0] * small.shape[1]:
        return None
    hull = cv2.convexHull(grand)
    peri = cv2.arcLength(hull, True)
    quad = None
    for eps in (0.01, 0.02, 0.03, 0.05, 0.08):
        approx = cv2.approxPolyDP(hull, eps * peri, True)
        if len(approx) == 4:
            quad = approx.reshape(4, 2).astype("float32")
            break
    if quad is None:
        quad = cv2.boxPoints(cv2.minAreaRect(hull)).astype("float32")
    return quad / k


def _text_mask(bgr):
    """
    Pixels de texte imprimé (sombres par rapport à leur voisinage), débarrassés des petites taches :
    sert à vérifier qu'une découpe ne coupe aucun texte, et à découper autour du contenu.
    Retourne (masque, échelle).
    """
    import cv2
    import numpy as np

    h, w = bgr.shape[:2]
    k = min(1.0, 1000.0 / max(h, w))
    small = cv2.resize(bgr, None, fx=k, fy=k, interpolation=cv2.INTER_AREA) if k < 1 else bgr
    g = cv2.cvtColor(small, cv2.COLOR_BGR2GRAY)
    fond = cv2.medianBlur(g, 31)
    T = (cv2.subtract(fond, g) > 25).astype(np.uint8) * 255
    s = max(3, int(0.015 * max(g.shape)) | 1)
    fusion = cv2.dilate(T, cv2.getStructuringElement(cv2.MORPH_RECT, (s * 3, s)))
    n, labels, stats, _ = cv2.connectedComponentsWithStats(fusion)
    garde = np.zeros_like(T)
    if n > 1:
        # on garde les blocs de texte importants ; les petits motifs (broderie, rayures du support) sont ignorés
        plus_grand = int(stats[1:, cv2.CC_STAT_AREA].max())
        seuil = max(0.002 * g.shape[0] * g.shape[1], 0.08 * plus_grand)
        for i in range(1, n):
            if stats[i, cv2.CC_STAT_AREA] >= seuil:
                garde[labels == i] = 255
    return cv2.bitwise_and(T, garde), k


def _couverture(T, quad_petit):
    """Part du texte qui se trouve à l'intérieur de la découpe (1.0 = rien n'est coupé)."""
    import cv2
    import numpy as np

    total = int((T > 0).sum())
    if total == 0:
        return 1.0
    masque = np.zeros_like(T)
    cv2.fillConvexPoly(masque, quad_petit.astype("int32"), 255)
    return float(((T > 0) & (masque > 0)).sum()) / total


def scan_document(img: Image.Image, decouper: bool = True) -> Image.Image:
    """
    « Numérisation » : détecte la feuille, la découpe, redresse la perspective, puis uniformise l'éclairage
    (fond blanc, texte net) SANS tout passer en noir et blanc brutal : le texte fin reste lisible.
    Si la feuille n'est pas trouvée, l'image entière est conservée (éclairage uniformisé seulement).
    """
    import cv2
    import numpy as np

    bgr = cv2.cvtColor(np.array(ImageOps.exif_transpose(img).convert("RGB")), cv2.COLOR_RGB2BGR)
    H0, W0 = bgr.shape[:2]
    T, kt = _text_mask(bgr)
    page = _page_quad(bgr) if decouper else None
    if page is not None and T.any() and _couverture(T, page * kt) < 0.97:
        page = None                       # la découpe « par la couleur » couperait du texte : on l'abandonne
    if decouper and page is None and T.any():
        # secours : on découpe autour du CONTENU imprimé (marge de 2 %), sans jamais couper de texte
        x, y, wd, ht = cv2.boundingRect(cv2.findNonZero(T))
        m = 0.02 * max(wd, ht)
        x0, y0 = max(0, (x - m) / kt), max(0, (y - m) / kt)
        x1, y1 = min(W0, (x + wd + m) / kt), min(H0, (y + ht + m) / kt)
        if (x1 - x0) * (y1 - y0) < 0.92 * W0 * H0:
            page = np.array([[x0, y0], [x1, y0], [x1, y1], [x0, y1]], dtype="float32")
    if page is not None:
        tl, tr, br, bl = _order_points(page)
        wd = int(max(np.linalg.norm(br - bl), np.linalg.norm(tr - tl)))
        ht = int(max(np.linalg.norm(tr - br), np.linalg.norm(tl - bl)))
        if wd > 200 and ht > 200:
            ech = min(1.0, 3000.0 / max(wd, ht))
            wd, ht = int(wd * ech), int(ht * ech)
            dst = np.array([[0, 0], [wd - 1, 0], [wd - 1, ht - 1], [0, ht - 1]], dtype="float32")
            M = cv2.getPerspectiveTransform(np.array([tl, tr, br, bl], dtype="float32"), dst)
            bgr = cv2.warpPerspective(bgr, M, (wd, ht), flags=cv2.INTER_CUBIC)
            marge = max(2, int(0.006 * max(wd, ht)))          # petite marge : retire le bord sombre résiduel
            bgr = bgr[marge:-marge, marge:-marge]

    gris = cv2.cvtColor(bgr, cv2.COLOR_BGR2GRAY)
    # éclairage uniformisé : on divise par le fond estimé (ombres, dégradés) puis on étire le contraste
    h, w = gris.shape
    petit = cv2.resize(gris, None, fx=0.25, fy=0.25, interpolation=cv2.INTER_AREA)
    k = max(3, (min(petit.shape) // 12) | 1)
    fond = cv2.resize(cv2.medianBlur(petit, k), (w, h), interpolation=cv2.INTER_LINEAR)
    norm_ = cv2.divide(gris, np.maximum(fond, 1), scale=255)
    bas, haut = np.percentile(norm_, (1, 97))
    norm_ = np.clip((norm_.astype("float32") - bas) * 255.0 / max(haut - bas, 1.0), 0, 255).astype("uint8")
    return Image.fromarray(norm_).convert("RGB")


def image_quality(img: Image.Image) -> dict:
    """Netteté et taille de la photo : permet de prévenir « photo floue » avant de lire la facture."""
    import cv2
    import numpy as np

    g = cv2.cvtColor(np.array(img.convert("RGB")), cv2.COLOR_RGB2GRAY)
    h, w = g.shape
    k = 1000.0 / max(h, w)
    g2 = cv2.resize(g, None, fx=k, fy=k, interpolation=cv2.INTER_AREA)
    nettete = float(cv2.Laplacian(g2, cv2.CV_64F).var())
    return {"largeur": w, "hauteur": h, "nettete": round(nettete, 1),
            "floue": nettete < SEUIL_NETTETE or max(w, h) < 1000}


SEUIL_NETTETE = 150.0


def preprocess_for_ocr(img: Image.Image) -> Image.Image:
    """Niveaux de gris, agrandi vers ~2400 px sur le grand côté (le texte fin devient lisible pour l'OCR)."""
    img = ImageOps.exif_transpose(img).convert("L")
    w, h = img.size
    longest = max(w, h)
    cible = 2400
    if longest < cible:
        k = min(cible / longest, 3.0)
        img = img.resize((int(w * k), int(h * k)), Image.LANCZOS)
    elif longest > 3500:
        k = 3500 / longest
        img = img.resize((int(w * k), int(h * k)), Image.LANCZOS)
    return ImageOps.autocontrast(img, cutoff=1)


def _tesseract():
    try:
        import pytesseract
    except ImportError as e:
        raise OcrUnavailable("Installez pytesseract : pip install pytesseract") from e
    if os.environ.get("TESSERACT_CMD"):
        pytesseract.pytesseract.tesseract_cmd = os.environ["TESSERACT_CMD"]
    return pytesseract


def ocr_tesseract(img: Image.Image, psm: int = 6) -> str:
    pytesseract = _tesseract()
    img = preprocess_for_ocr(img)
    try:
        try:
            return pytesseract.image_to_string(img, lang="fra+eng", config=f"--psm {psm}")
        except pytesseract.TesseractError:      # pack de langue français absent
            return pytesseract.image_to_string(img, lang="eng", config=f"--psm {psm}")
    except pytesseract.TesseractNotFoundError as e:
        raise OcrUnavailable("Le programme Tesseract n'est pas installé (ou pas dans le PATH). "
                             "Voir les instructions d'installation.") from e


def find_total_in_image(img: Image.Image):
    """
    Montant « Net à payer / Total TTC » repéré par sa POSITION : à droite du libellé, ou juste en dessous
    (cas des tableaux où le libellé est en en-tête de colonne). Retourne un float ou None.
    """
    pytesseract = _tesseract()
    prepared = preprocess_for_ocr(img)
    try:
        d = pytesseract.image_to_data(prepared, lang="fra+eng", config="--psm 6",
                                      output_type=pytesseract.Output.DICT)
    except Exception:  # noqa: BLE001
        return None
    mots = []
    for i, t in enumerate(d["text"]):
        t = t.strip()
        if not t:
            continue
        mots.append({"t": t, "n": norm(t), "x": d["left"][i], "y": d["top"][i], "w": d["width"][i], "h": d["height"][i]})
    for m in mots:
        m["cx"], m["cy"], m["x2"] = m["x"] + m["w"] / 2, m["y"] + m["h"] / 2, m["x"] + m["w"]
    montants = []
    for m in mots:
        t = m["t"].replace("€", "").replace("EUR", "").strip()
        mt = _AMOUNT.fullmatch(t)
        if mt and 0 < _to_float(mt) < 100_000:
            montants.append((m, _to_float(mt)))
    meilleurs = []                                           # (priorité, distance, montant)
    for k in (m for m in mots if m["n"].startswith("ttc")):
        ligne = sorted((m for m in mots if abs(m["cy"] - k["cy"]) < 0.6 * k["h"]), key=lambda m: m["x"])
        i = ligne.index(k)
        gauche, droite = i, i
        sep = re.compile(r"^[\W_]+$")               # tirets, barres... : séparent les colonnes d'un en-tête
        # un libellé = des mots serrés les uns contre les autres ; un espace large ou un tiret = autre colonne
        while (gauche > 0 and ligne[gauche]["x"] - ligne[gauche - 1]["x2"] < 1.1 * k["h"]
               and not sep.match(ligne[gauche - 1]["t"])):
            gauche -= 1
        while (droite < len(ligne) - 1 and ligne[droite + 1]["x"] - ligne[droite]["x2"] < 1.1 * k["h"]
               and not sep.match(ligne[droite + 1]["t"])):
            droite += 1
        libelle = norm(" ".join(m["t"] for m in ligne[gauche:droite + 1]))
        payer = bool(re.search(r"payer", libelle))                       # « net à payer », « total à payer »...
        if not (payer or re.search(r"total|montant|net", libelle)):
            continue
        prio = 0 if payer else 1
        x0, x1 = ligne[gauche]["x"], ligne[droite]["x2"]
        marge = 0.25 * (x1 - x0)
        for m, valeur in montants:
            if abs(m["cy"] - k["cy"]) < 0.6 * k["h"] and m["x"] >= x1 - 0.5 * k["h"] and m["x"] - x1 < 20 * k["h"]:
                meilleurs.append((prio, 0, m["x"] - x1, valeur))                  # même ligne, à droite
            elif payer and 0.5 * k["h"] < m["cy"] - k["cy"] < 12 * k["h"] and x0 - marge <= m["cx"] <= x1 + marge:
                # juste sous « Net à payer TTC » (libellé en en-tête de colonne). On ne le fait que pour « payer » :
                # « Total TTC » en tête d'un tableau d'articles est suivi des prix de chaque ligne, pas du total.
                meilleurs.append((prio, 1, m["cy"] - k["cy"], valeur))
    if not meilleurs:
        return None
    return sorted(meilleurs)[0][3]


_EASY = None


def ocr_easyocr(img: Image.Image) -> str:
    global _EASY
    try:
        import easyocr
        import numpy as np
    except ImportError as e:
        raise OcrUnavailable("Installez EasyOCR : pip install easyocr") from e
    if _EASY is None:
        _EASY = easyocr.Reader(["fr", "en"], gpu=False)
    arr = np.array(preprocess_for_ocr(img).convert("RGB"))
    return "\n".join(_EASY.readtext(arr, detail=0, paragraph=False))


def run_ocr(images, engine: str = "Tesseract", psm: int = 6) -> str:
    if engine.lower().startswith("easy"):
        return "\n".join(ocr_easyocr(im) for im in images)
    return "\n".join(ocr_tesseract(im, psm) for im in images)


def read_invoice(images, engine: str = "Tesseract") -> dict:
    """
    Lit la facture : texte, montant TTC le plus fiable et montants candidats.
    Passe 1 : lecture normale. Passe 2 (si aucun montant) : autre découpage de la page.
    Le montant repéré par sa position (sous/à côté de « Net à payer TTC ») prime sur l'analyse du texte.
    """
    texte = run_ocr(images, engine)
    meilleur, candidats = find_amount(texte)
    total = None
    if not engine.lower().startswith("easy"):
        for im in images:
            total = find_total_in_image(im)
            if total:
                break
        if meilleur is None and total is None:
            texte += "\n" + run_ocr(images, engine, psm=4)
            meilleur, candidats = find_amount(texte)
    somme = _total_par_somme([float(m.group(0).replace(",", ".")) for m in re.finditer(r"\d+[.,]\d{2}", texte)])
    if total and somme and total != somme[0] and total in somme[1]:
        total = somme[0]                    # ex. 3,88 (TVA) repéré à tort : 19,42 + 3,88 = 23,30
    if total:
        meilleur = total
        candidats = [total] + [c for c in candidats if c != total]
    return {"text": texte, "amount": meilleur, "candidates": candidats[:6]}


# ════════════════════════════════════════════════════════════════════════════
#  5. EXTRACTION : montant TTC, date, catégorie
# ════════════════════════════════════════════════════════════════════════════
_AMOUNT = re.compile(
    r"(?<![\d.,/\-])(\d{1,3}(?:[ \u00a0.]\d{3})+|\d+)\s?[.,]\s?(\d{2})"
    r"(?!\d)(?![.,/\-]\d)(?!\s?%)"
)


def _to_float(m: re.Match) -> float:
    ent = re.sub(r"[ \u00a0.]", "", m.group(1))
    return float(f"{ent}.{m.group(2)}")


def _line_weight(n: str) -> int:
    """3 = ligne « TTC / net à payer », 2 = « total / montant / CB », 0 = à ignorer."""
    if "sous-total" in n or "sous total" in n:
        return 0
    if re.search(r"net a payer|total ttc|montant ttc|total a payer|montant a payer|"
                 r"a payer ttc|total general|total facture", n) or re.search(r"\bttc\b", n):
        return 3
    if re.search(r"\b(ht|tva|hors taxes?)\b", n):
        return 0
    if re.search(r"\b(total|montant|somme|carte|cb|especes|paye|reglement)\b", n):
        return 2
    return 0


def _total_par_somme(valeurs):
    """
    Cherche un montant égal à la somme de deux autres (HT + TVA = TTC, ou somme des lignes).
    Retourne (total, (a, b)) pour le plus grand total trouvé, ou None.
    """
    vals = sorted(round(v, 2) for v in valeurs if 0 < v < 100_000)
    meilleur = None
    for i, a in enumerate(vals):
        for b in vals[i + 1:]:
            c = round(a + b, 2)
            if c > b and c in vals and (meilleur is None or c > meilleur[0]):
                meilleur = (c, (a, b))
    return meilleur


def find_amount(text: str):
    """Retourne (meilleur montant TTC ou None, liste des montants candidats)."""
    lines = text.splitlines()
    weighted, every = [], []
    for i, line in enumerate(lines):
        found = [_to_float(m) for m in _AMOUNT.finditer(line)]
        every += found
        n_ligne = norm(line)
        w = _line_weight(n_ligne)
        entete = bool(re.search(r"\b(ht|tva)\b", n_ligne))          # en-tête de colonnes : la ligne suivante n'est pas fiable
        if w and not found and i + 1 < len(lines) and not entete:     # montant sur la ligne suivante
            found = [_to_float(m) for m in _AMOUNT.finditer(lines[i + 1])]
        weighted += [(w, a) for a in found if w]
    every = [a for a in every if 0 < a < 100_000]
    weighted = [(w, a) for w, a in weighted if 0 < a < 100_000]

    if weighted:
        best = max(weighted)[1]
    else:
        somme = _total_par_somme(every)
        best = somme[0] if somme else (max(every) if every else None)
    cands = []
    for a in ([best] if best else []) + [a for _, a in sorted(weighted, reverse=True)] + sorted(every, reverse=True):
        if a not in cands:
            cands.append(a)
    return best, cands[:6]


_MOIS_PREFIX = {"janv": 1, "jan": 1, "fevr": 2, "fev": 2, "mars": 3, "avr": 4, "mai": 5, "juin": 6,
                "juil": 7, "aout": 8, "sept": 9, "sep": 9, "oct": 10, "nov": 11, "dec": 12}
_D_NUM = re.compile(r"(?<!\d)(\d{1,2})\s?[/\-.]\s?(\d{1,2})\s?[/\-.]\s?(\d{4}|\d{2})(?!\d)")
_D_ISO = re.compile(r"(?<!\d)(\d{4})-(\d{2})-(\d{2})(?!\d)")
_D_TXT = re.compile(r"(?<!\d)(\d{1,2})(?:er)?\s+(janv|jan|fevr|fev|mars|avr|mai|juin|juil|aout|sept|sep|oct|nov|dec)"
                    r"[a-z]*\.?\s+(\d{4}|\d{2})(?!\d)")


def find_date(text: str, today: dt.date | None = None) -> dt.date | None:
    today = today or dt.date.today()
    best, best_score = None, None
    for line in text.splitlines():
        n = norm(line)
        cands = []
        for m in _D_NUM.finditer(n):
            cands.append((int(m.group(3)), int(m.group(2)), int(m.group(1)), len(m.group(3)) == 4))
        for m in _D_ISO.finditer(n):
            cands.append((int(m.group(1)), int(m.group(2)), int(m.group(3)), True))
        for m in _D_TXT.finditer(n):
            cands.append((int(m.group(3)), _MOIS_PREFIX[m.group(2)], int(m.group(1)), len(m.group(3)) == 4))
        for y, mo, d, four in cands:
            y = y + 2000 if y < 100 else y
            try:
                date = dt.date(y, mo, d)
            except ValueError:
                continue
            if not (2000 <= y <= today.year + 1):
                continue
            score = (2 if re.search(r"\bdate\b", n) and not re.search(r"echeance|limite|naissance", n) else 0) \
                - (2 if re.search(r"echeance|limite|avant le", n) else 0) + (1 if four else 0)
            if best_score is None or score > best_score:       # à score égal : la première trouvée
                best, best_score = date, score
    return best


_GENERIC = {"achat", "achats", "divers", "diverses", "fourniture", "fournitures", "frais", "total"}
_SYNONYMES = {"carburant": ["essence", "gazole", "diesel", "sans plomb", "sp95", "sp98", "e10"],
              "assurance": ["assurance", "cotisation assurance", "attestation"],
              "bureau": ["bureautique", "papeterie", "cartouche", "ramette", "stylo", "classeur", "agrafe"],
              "fumier": ["fumier", "compost"]}
# Lignes qui décrivent seulement le MOYEN DE PAIEMENT : elles ne disent rien sur la nature de l'achat
# (« carte bancaire » figure sur presque tous les tickets de caisse).
_PAIEMENT = re.compile(r"carte (bancaire|bleue|banc)|\b(cb|visa|mastercard|espece|especes|paiement|reglement|"
                       r"sans contact|contactless|monnaie rendue)\b")


def suggest_category(text: str, labels: list[str]) -> int | None:
    """Index de la catégorie la plus plausible d'après la nature de l'achat (ou None si aucun indice net)."""
    lignes = [l for l in text.splitlines() if not _PAIEMENT.search(norm(l))]
    t = norm(" ".join(lignes))
    best_i, best = None, 0
    for i, label in enumerate(labels):
        score = 0
        for tok in re.findall(r"[a-z0-9]{4,}", norm(label)):
            if tok in _GENERIC:
                continue
            if tok in t:
                score += 2
            score += sum(1 for syn in _SYNONYMES.get(tok, []) if syn in t)
        if score > best:
            best_i, best = i, score
    return best_i


def prepare_justificatif(data: bytes, filename: str, preview: Image.Image, max_pdf: int = 4_000_000):
    """
    Pièce justificative à conserver : (octets, type MIME).
    PDF d'origine s'il est léger, sinon image JPEG réduite (1600 px, qualité 80 : environ 300 Ko).
    """
    is_pdf = filename.lower().endswith(".pdf") or data[:5] == b"%PDF-"
    if is_pdf and len(data) <= max_pdf:
        return data, "application/pdf"
    img = ImageOps.exif_transpose(preview).convert("RGB")
    img.thumbnail((1600, 1600))
    out = io.BytesIO()
    img.save(out, "JPEG", quality=80, optimize=True)
    return out.getvalue(), "image/jpeg"


def excel_verrouille(path) -> bool:
    """
    Vrai si le fichier semble ouvert dans Excel (Windows interdit alors d'y écrire) :
    fichier verrou « ~$nom.xlsx » présent, ou écriture refusée par le système.
    """
    p = Path(path)
    for nom in ("~$" + p.name, "~$" + p.name[2:]):
        if (p.parent / nom).exists():
            return True
    try:
        with open(p, "ab"):
            pass
    except PermissionError:
        return True
    except OSError:
        return False
    return False
